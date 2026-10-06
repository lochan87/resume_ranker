"""Resume ingestion module: discovery, format parsing, link extraction, and deduplication."""

import hashlib
import logging
from pathlib import Path
from typing import Optional
import fitz  # PyMuPDF
import pdfplumber
import docx

from src.models import ParsedResume, IngestionFailure

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def compute_file_hash(path: Path) -> str:
    """Compute SHA-256 hash of file content."""
    hasher = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def extract_pdf_content(path: Path) -> tuple[str, list[str]]:
    """Extract text and hyperlinks from PDF using PyMuPDF, falling back to pdfplumber."""
    text_chunks: list[str] = []
    links: list[str] = []

    try:
        doc = fitz.open(path)
        if doc.is_encrypted:
            if not doc.authenticate(""):
                raise ValueError("Encrypted PDF document cannot be read")

        for page in doc:
            page_text = page.get_text()
            if page_text and page_text.strip():
                text_chunks.append(page_text.strip())

            # Extract hyperlink annotations
            for link in page.get_links():
                uri = link.get("uri")
                if uri and uri.strip():
                    links.append(uri.strip())
        doc.close()
    except Exception as exc:
        logger.warning(f"PyMuPDF failed on {path.name}: {exc}. Trying pdfplumber fallback.")
        text_chunks.clear()

    # Fallback to pdfplumber if PyMuPDF returned no text or errored
    if not text_chunks:
        try:
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    p_text = page.extract_text()
                    if p_text and p_text.strip():
                        text_chunks.append(p_text.strip())
        except Exception as exc:
            raise ValueError(f"Failed to read PDF with pdfplumber fallback: {exc}") from exc

    raw_text = "\n\n".join(text_chunks).strip()
    if not raw_text:
        raise ValueError("Empty or scanned-image PDF without extractable text")

    # Append extracted hyperlink URLs to raw_text
    if links:
        raw_text += "\n\nExtracted Links:\n" + "\n".join(links)

    return raw_text, links


def extract_docx_content(path: Path) -> tuple[str, list[str]]:
    """Extract text from DOCX document."""
    try:
        doc = docx.Document(path)
        chunks: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_str:
                    chunks.append(row_str)
        text = "\n\n".join(chunks).strip()
        if not text:
            raise ValueError("DOCX document has no text content")
        return text, []
    except Exception as exc:
        raise ValueError(f"Failed to parse DOCX: {exc}") from exc


def extract_txt_content(path: Path) -> tuple[str, list[str]]:
    """Extract text from TXT file."""
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read().strip()
        if not text:
            raise ValueError("TXT file is empty")
        return text, []
    except Exception as exc:
        raise ValueError(f"Failed to read TXT file: {exc}") from exc


def parse_resume_file(path: Path) -> tuple[str, list[str]]:
    """Route file parsing by extension."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf_content(path)
    elif ext == ".docx":
        return extract_docx_content(path)
    elif ext == ".txt":
        return extract_txt_content(path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")


def discover_resume_files(directory: Path | str) -> list[Path]:
    """Recursively discover supported resume files."""
    base = Path(directory)
    if not base.exists() or not base.is_dir():
        return []
    files: list[Path] = [
        p for p in base.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    return sorted(files)


def ingest_file(
    path: Path, seen_hashes: dict[str, str]
) -> tuple[Optional[ParsedResume], Optional[IngestionFailure]]:
    """Ingest a single resume file with hash-based deduplication and error isolation."""
    try:
        file_hash = compute_file_hash(path)
        if file_hash in seen_hashes:
            orig = seen_hashes[file_hash]
            return (
                ParsedResume(
                    source_file=str(path),
                    content_hash=file_hash,
                    raw_text="",
                    is_duplicate=True,
                    duplicate_of=orig,
                ),
                None,
            )

        text, links = parse_resume_file(path)
        seen_hashes[file_hash] = str(path)
        return (
            ParsedResume(
                source_file=str(path),
                content_hash=file_hash,
                raw_text=text,
                links=links,
                is_duplicate=False,
            ),
            None,
        )
    except Exception as exc:
        return None, IngestionFailure(file=str(path), reason=str(exc))


def ingest_resumes(
    directory_or_files: Path | str | list[Path],
) -> tuple[list[ParsedResume], list[IngestionFailure]]:
    """Ingest multiple resumes, returning parsed resumes and failures."""
    if isinstance(directory_or_files, list):
        files = directory_or_files
    else:
        files = discover_resume_files(directory_or_files)

    parsed_list: list[ParsedResume] = []
    failures: list[IngestionFailure] = []
    seen_hashes: dict[str, str] = {}

    for file_path in files:
        parsed, failure = ingest_file(file_path, seen_hashes)
        if failure:
            failures.append(failure)
        elif parsed:
            parsed_list.append(parsed)

    return parsed_list, failures
